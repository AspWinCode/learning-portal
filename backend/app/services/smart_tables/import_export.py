"""Import/export CSV и XLSX — Phase 4.

Импорт создаёт НОВЫЙ лист в workbook (не трогает существующие данные —
см. docs/smart-tables-architecture.md, раздел M: "Импорт CSV 10k строк
создаёт sheet с нужными типами колонок"), поэтому идёт в обход
OperationExecutor — это не "операция" над существующим состоянием, а
объёмная массовая вставка, для которой undo через operation log был бы
неоправданно дорог (пришлось бы хранить весь импортированный датасет
дважды). Экспорт отдаёт текущий computed_value каждой ячейки (не
формулу — значение).
"""
from __future__ import annotations

import csv
import io
import re
from datetime import date, datetime
from typing import List, Optional, Sequence, Tuple

from openpyxl import Workbook as XlsxWorkbook
from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.models import SmartTableCell, SmartTableColumn, SmartTableRow, SmartTableSheet, SmartTableWorkbook
from app.services.smart_tables.executor import _coerce_value

_DATE_RE_ISO = re.compile(r"^\d{4}-\d{2}-\d{2}$")
_DATE_RE_DMY = re.compile(r"^\d{1,2}[./]\d{1,2}[./]\d{2,4}$")
_BOOL_TRUE = {"true", "да", "yes", "истина"}
_BOOL_FALSE = {"false", "нет", "no", "ложь"}


class ImportError_(ValueError):
    """Файл нельзя разобрать (пустой, битый, неподдерживаемый формат)."""


# ── разбор файла в таблицу строк ───────────────────────────────

def parse_upload(filename: str, data: bytes) -> Tuple[List[str], List[List[Optional[str]]]]:
    name = (filename or "").lower()
    if name.endswith(".csv"):
        return _parse_csv(data)
    if name.endswith(".xlsx"):
        return _parse_xlsx(data)
    raise ImportError_("Поддерживаются только файлы .csv и .xlsx")


def _parse_csv(data: bytes) -> Tuple[List[str], List[List[Optional[str]]]]:
    text = data.decode("utf-8-sig", errors="replace")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    reader = csv.reader(io.StringIO(text), dialect)
    rows = [row for row in reader]
    if not rows:
        raise ImportError_("Пустой CSV-файл")
    header, *body = rows
    return [h.strip() for h in header], [[(c if c != "" else None) for c in r] for r in body]


def _parse_xlsx(data: bytes) -> Tuple[List[str], List[List[Optional[str]]]]:
    wb = load_workbook(filename=io.BytesIO(data), data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise ImportError_("Пустой лист в XLSX")
    header, *body = rows
    headers = [str(h).strip() if h is not None else f"Колонка {i + 1}" for i, h in enumerate(header)]
    out_rows = []
    for r in body:
        out_rows.append([_xlsx_cell_to_str(v) for v in r])
    return headers, out_rows


def _xlsx_cell_to_str(v) -> Optional[str]:
    if v is None:
        return None
    if isinstance(v, (datetime, date)):
        return v.isoformat()
    return str(v)


# ── автоопределение типа колонки ───────────────────────────────

def _looks_number(v: str) -> bool:
    try:
        float(v.replace(",", ".").replace(" ", ""))
        return True
    except ValueError:
        return False


def _looks_boolean(v: str) -> bool:
    return v.strip().lower() in _BOOL_TRUE or v.strip().lower() in _BOOL_FALSE


def _looks_date(v: str) -> bool:
    return bool(_DATE_RE_ISO.match(v.strip()) or _DATE_RE_DMY.match(v.strip()))


def infer_column_type(values: Sequence[Optional[str]]) -> str:
    non_empty = [v for v in values if v is not None and str(v).strip() != ""]
    if not non_empty:
        return "text"
    if all(_looks_boolean(v) for v in non_empty):
        return "boolean"
    if all(_looks_number(v) for v in non_empty):
        return "number"
    if all(_looks_date(v) for v in non_empty):
        return "date"
    return "text"


# ── импорт: создаёт новый лист ─────────────────────────────────

def import_sheet(
    db: Session,
    workbook: SmartTableWorkbook,
    sheet_name: str,
    headers: List[str],
    rows: List[List[Optional[str]]],
    *,
    max_rows: int = 20000,
) -> SmartTableSheet:
    if len(rows) > max_rows:
        raise ImportError_(f"Слишком много строк ({len(rows)}) — максимум {max_rows} за один импорт")
    if not headers:
        raise ImportError_("В файле нет заголовков колонок")

    existing = db.query(SmartTableSheet.position).filter(SmartTableSheet.workbook_id == workbook.id).all()
    position = (max((p for (p,) in existing), default=-1.0)) + 1.0

    sheet = SmartTableSheet(workbook_id=workbook.id, name=sheet_name, position=position)
    db.add(sheet)
    db.flush()

    columns: List[SmartTableColumn] = []
    for i, name in enumerate(headers):
        col_values = [r[i] if i < len(r) else None for r in rows]
        col_type = infer_column_type(col_values)
        column = SmartTableColumn(
            sheet_id=sheet.id, name=name or f"Колонка {i + 1}", position=float(i), type=col_type, config={},
        )
        db.add(column)
        columns.append(column)
    db.flush()

    for row_idx, raw_row in enumerate(rows):
        row = SmartTableRow(sheet_id=sheet.id, position=float(row_idx), cells_snapshot={})
        db.add(row)
        db.flush()
        snapshot = {}
        for i, column in enumerate(columns):
            raw_value = raw_row[i] if i < len(raw_row) else None
            if raw_value is None or str(raw_value).strip() == "":
                continue
            computed = _coerce_value(raw_value, column.type)
            db.add(SmartTableCell(
                row_id=row.id, column_id=column.id, raw_value=str(raw_value), formula=None,
                computed_value=computed, formatting={},
            ))
            snapshot[str(column.id)] = {"value": computed, "formula": None, "formatting": {}, "metadata": {}}
        row.cells_snapshot = snapshot

    db.flush()
    return sheet


# ── экспорт ─────────────────────────────────────────────────────

def _display(value) -> str:
    if value is None:
        return ""
    if isinstance(value, dict) and "error" in value:
        return value["error"]
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    return str(value)


def export_csv(sheet: SmartTableSheet, columns: List[SmartTableColumn], rows: List[SmartTableRow]) -> bytes:
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow([c.name for c in columns])
    for row in rows:
        snapshot = row.cells_snapshot or {}
        writer.writerow([_display((snapshot.get(str(c.id)) or {}).get("value")) for c in columns])
    return buf.getvalue().encode("utf-8-sig")


def export_xlsx(sheet: SmartTableSheet, columns: List[SmartTableColumn], rows: List[SmartTableRow]) -> bytes:
    wb = XlsxWorkbook()
    ws = wb.active
    ws.title = (sheet.name or "Sheet")[:31] or "Sheet"
    ws.append([c.name for c in columns])
    for row in rows:
        snapshot = row.cells_snapshot or {}
        values = []
        for c in columns:
            v = (snapshot.get(str(c.id)) or {}).get("value")
            if isinstance(v, dict) and "error" in v:
                v = v["error"]
            values.append(v)
        ws.append(values)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
